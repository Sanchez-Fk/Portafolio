#!/usr/bin/env python3
"""
Importa leads desde un CSV o un Excel a la base `prospector`.

    python importar_leads.py mis-leads.csv
    python importar_leads.py leads.xlsx --hoja "Restaurantes" --origen "google maps"
    python importar_leads.py leads.csv --simular      # no escribe nada

Columnas que entiende (en español o inglés, con o sin tildes, en cualquier orden):

    telefono*  nombre  negocio*  rubro  ciudad  sitio_web  observacion  origen

    * obligatorias. Si no hay columna `negocio`, se usa `nombre`.

Cualquier columna que no reconozca no se pierde: se guarda en el campo `extra`
(JSON), por si después la quieres usar en el mensaje.

Lo que hace con los teléfonos:
  - deja solo dígitos,
  - quita el 00 o el + inicial,
  - si el número viene sin código de país, le pone --pais (57 por defecto),
  - descarta lo que no pueda quedar entre 10 y 15 dígitos.

Los repetidos no se duplican: un teléfono que ya está en la base se ACTUALIZA
(rellena los campos vacíos) pero nunca vuelve al estado `nuevo`. Es decir, no
puedes re-encolar sin querer a alguien al que ya le escribiste.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

try:
    import psycopg
except ImportError:
    sys.exit("Falta psycopg. Instálalo con:  pip install -r requirements.txt")


# ─── Nombres de columna aceptados ────────────────────────────────────────────

ALIAS = {
    "telefono":    {"telefono", "telefonos", "phone", "celular", "movil", "whatsapp",
                    "numero", "num", "tel", "contacto_telefono"},
    "nombre":      {"nombre", "name", "contacto", "propietario", "dueno", "encargado",
                    "nombre_contacto", "first_name"},
    "negocio":     {"negocio", "empresa", "business", "company", "razon_social",
                    "nombre_negocio", "establecimiento", "local"},
    "rubro":       {"rubro", "categoria", "category", "giro", "sector", "tipo",
                    "industria", "industry"},
    "ciudad":      {"ciudad", "city", "municipio", "localidad", "zona"},
    "sitio_web":   {"sitio_web", "web", "website", "url", "pagina", "pagina_web", "sitio"},
    "observacion": {"observacion", "observaciones", "nota", "notas", "detalle",
                    "comentario", "que_le_falta", "dolor"},
    "origen":      {"origen", "fuente", "source", "de_donde"},
}


def normalizar_cabecera(texto: str) -> str:
    """"Sitio Web " → "sitio_web".  "Teléfono" → "telefono"."""
    limpio = unicodedata.normalize("NFKD", str(texto or "").strip().lower())
    limpio = "".join(c for c in limpio if not unicodedata.combining(c))
    limpio = re.sub(r"[^a-z0-9]+", "_", limpio).strip("_")
    return limpio


def mapear_columnas(cabeceras: list[str]) -> dict[str, str]:
    """Devuelve {campo_del_esquema: cabecera_original}."""
    mapa: dict[str, str] = {}
    for cabecera in cabeceras:
        clave = normalizar_cabecera(cabecera)
        for campo, alias in ALIAS.items():
            if campo not in mapa and clave in alias:
                mapa[campo] = cabecera
                break
    return mapa


# ─── Teléfonos ───────────────────────────────────────────────────────────────

# Largos nacionales típicos en LATAM: 10 (CO, MX, AR sin 0), 8-9 (CL, PE, EC…).
LARGO_NACIONAL = range(8, 11)


def normalizar_telefono(bruto, pais: str) -> str | None:
    """Devuelve E.164 sin '+', o None si el número no es rescatable."""
    if bruto is None:
        return None

    texto = str(bruto).strip()
    # Excel guarda los teléfonos como float: 573052624583.0
    if re.fullmatch(r"\d+\.0", texto):
        texto = texto[:-2]

    digitos = re.sub(r"\D", "", texto)
    if not digitos:
        return None

    digitos = re.sub(r"^00", "", digitos)

    # ¿Ya trae código de país?
    if digitos.startswith(pais) and len(digitos) > len(pais) + 7:
        completo = digitos
    elif len(digitos) in LARGO_NACIONAL:
        completo = pais + digitos
    elif len(digitos) > 15:
        # Dos números pegados en la misma celda, o basura.
        return None
    else:
        completo = digitos

    return completo if 10 <= len(completo) <= 15 else None


# ─── Lectura de archivos ─────────────────────────────────────────────────────

def leer_csv(ruta: Path) -> list[dict]:
    # utf-8-sig se come el BOM que mete Excel al exportar.
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            with ruta.open(encoding=codificacion, newline="") as f:
                muestra = f.read(4096)
                f.seek(0)
                try:
                    dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
                except csv.Error:
                    dialecto = csv.excel
                return [dict(fila) for fila in csv.DictReader(f, dialect=dialecto)]
        except UnicodeDecodeError:
            continue
    sys.exit(f"No pude leer {ruta}: codificación desconocida.")


def leer_excel(ruta: Path, hoja: str | None) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Para leer Excel falta openpyxl:  pip install -r requirements.txt")

    libro = load_workbook(ruta, read_only=True, data_only=True)
    pagina = libro[hoja] if hoja else libro[libro.sheetnames[0]]

    filas = pagina.iter_rows(values_only=True)
    try:
        cabeceras = [str(c) if c is not None else "" for c in next(filas)]
    except StopIteration:
        return []

    return [
        dict(zip(cabeceras, fila))
        for fila in filas
        if any(c is not None and str(c).strip() for c in fila)
    ]


# ─── Programa ────────────────────────────────────────────────────────────────

def main() -> int:
    p = argparse.ArgumentParser(description="Importa leads a la base prospector.")
    p.add_argument("archivo", type=Path, help="CSV, XLSX o XLSM")
    p.add_argument("--hoja", help="Hoja del Excel (por defecto, la primera)")
    p.add_argument("--pais", default=os.environ.get("PAIS_POR_DEFECTO", "57"),
                   help="Código de país para números sin prefijo (57 = Colombia)")
    p.add_argument("--origen", help="Etiqueta de origen para todo el archivo")
    p.add_argument("--dsn", default=os.environ.get("DATABASE_URL"),
                   help="postgresql://usuario:clave@host:5432/prospector")
    p.add_argument("--simular", action="store_true",
                   help="Analiza y reporta, pero no escribe en la base")
    args = p.parse_args()

    if not args.archivo.exists():
        sys.exit(f"No existe: {args.archivo}")
    if not args.simular and not args.dsn:
        sys.exit("Falta la conexión. Usa --dsn o exporta DATABASE_URL.")

    sufijo = args.archivo.suffix.lower()
    crudas = (leer_excel(args.archivo, args.hoja) if sufijo in {".xlsx", ".xlsm"}
              else leer_csv(args.archivo))

    if not crudas:
        sys.exit("El archivo no tiene filas.")

    mapa = mapear_columnas(list(crudas[0].keys()))
    if "telefono" not in mapa:
        sys.exit(
            "No encontré la columna del teléfono.\n"
            f"Columnas del archivo: {', '.join(str(c) for c in crudas[0])}\n"
            f"Nómbrala como alguna de: {', '.join(sorted(ALIAS['telefono']))}"
        )

    reconocidas = set(mapa.values())
    listos: dict[str, dict] = {}   # telefono → lead, ya deduplicado
    descartados: list[tuple[int, str, str]] = []

    for i, fila in enumerate(crudas, start=2):   # 2 = primera fila de datos en la hoja
        telefono = normalizar_telefono(fila.get(mapa["telefono"]), args.pais)
        if not telefono:
            descartados.append((i, str(fila.get(mapa["telefono"], "")), "teléfono inválido"))
            continue
        if telefono in listos:
            descartados.append((i, telefono, "repetido dentro del archivo"))
            continue

        def valor(campo: str) -> str | None:
            if campo not in mapa:
                return None
            v = fila.get(mapa[campo])
            v = str(v).strip() if v is not None else ""
            return v or None

        negocio = valor("negocio") or valor("nombre")
        if not negocio:
            descartados.append((i, telefono, "sin negocio ni nombre"))
            continue

        extra = {
            normalizar_cabecera(k): str(v).strip()
            for k, v in fila.items()
            if k not in reconocidas and v is not None and str(v).strip()
        }

        listos[telefono] = {
            "telefono": telefono,
            "nombre": valor("nombre"),
            "negocio": negocio,
            "rubro": valor("rubro"),
            "ciudad": valor("ciudad"),
            "sitio_web": valor("sitio_web"),
            "observacion": valor("observacion"),
            "origen": args.origen or valor("origen") or args.archivo.name,
            "extra": json.dumps(extra, ensure_ascii=False),
        }

    print(f"\nArchivo      {args.archivo.name}")
    print(f"Filas        {len(crudas)}")
    print(f"Válidos      {len(listos)}")
    print(f"Descartados  {len(descartados)}")
    print(f"Columnas     {', '.join(f'{k} ← {v}' for k, v in mapa.items())}")

    if descartados:
        print("\nDescartados (primeros 15):")
        for fila_n, dato, motivo in descartados[:15]:
            print(f"  fila {fila_n:>5}  {dato[:22]:<22}  {motivo}")

    if not listos:
        print("\nNada que importar.")
        return 1

    if args.simular:
        print("\n--simular: no se escribió nada. Ejemplo de lo que se insertaría:")
        primero = next(iter(listos.values()))
        print(json.dumps(primero, ensure_ascii=False, indent=2))
        return 0

    # `where leads.estado <> 'nuevo'` NO aparece en el update a propósito: los
    # campos se rellenan siempre, pero el estado y los contadores no se tocan
    # nunca. Reimportar el mismo archivo no reencola a nadie.
    sql = """
        insert into leads
          (telefono, nombre, negocio, rubro, ciudad, sitio_web, observacion, origen, extra)
        values
          (%(telefono)s, %(nombre)s, %(negocio)s, %(rubro)s, %(ciudad)s,
           %(sitio_web)s, %(observacion)s, %(origen)s, %(extra)s::jsonb)
        on conflict (telefono) do update set
          nombre      = coalesce(leads.nombre,      excluded.nombre),
          rubro       = coalesce(leads.rubro,       excluded.rubro),
          ciudad      = coalesce(leads.ciudad,      excluded.ciudad),
          sitio_web   = coalesce(leads.sitio_web,   excluded.sitio_web),
          observacion = coalesce(leads.observacion, excluded.observacion),
          extra       = leads.extra || excluded.extra
        returning (xmax = 0) as insertado
    """

    nuevos = actualizados = 0
    with psycopg.connect(args.dsn) as conexion:
        with conexion.cursor() as cur:
            for lead in listos.values():
                cur.execute(sql, lead)
                if cur.fetchone()[0]:
                    nuevos += 1
                else:
                    actualizados += 1

            cur.execute("select count(*) from bloqueos b join leads l using (telefono)")
            en_lista_negra = cur.fetchone()[0]
            cur.execute("select count(*) from leads where estado = 'nuevo'")
            en_cola = cur.fetchone()[0]

    print(f"\nInsertados     {nuevos}")
    print(f"Actualizados   {actualizados}  (ya estaban; no se reencolan)")
    if en_lista_negra:
        print(f"En lista negra {en_lista_negra}  (pidieron baja; nunca se les escribe)")
    print(f"\nEn cola ahora  {en_cola} leads en estado 'nuevo'.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
